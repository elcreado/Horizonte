import csv
import io
import re
from datetime import date, datetime, time
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook

from .imports import parse_csv


def parse_xlsx(content: bytes) -> list[dict]:
    """Lee datos, nunca fórmulas; limita ZIP y coordenadas antes de abrir el libro."""
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if (
                len(entries) > 200
                or len(set(names)) != len(names)
                or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024
            ):
                raise ValueError(
                    "XLSX excede el límite de 20 MB descomprimidos o contiene demasiadas partes."
                )
            for entry in entries:
                if (
                    "vbaproject" in entry.filename.lower()
                    or "externallinks" in entry.filename.lower()
                ):
                    raise ValueError("No se admiten macros ni enlaces a otros libros.")
                if entry.filename.endswith((".xml", ".rels")):
                    root = fromstring(
                        archive.read(entry),
                        forbid_dtd=True,
                        forbid_entities=True,
                        forbid_external=True,
                    )
                    if entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(
                        ".xml"
                    ):
                        for node in root.iter():
                            if (
                                node.tag.endswith("}row")
                                and not 1 <= int(node.get("r", "0")) <= 10001
                            ):
                                raise ValueError(
                                    "Máximo 10.000 movimientos y una fila de encabezados."
                                )
                            if node.tag.endswith("}c"):
                                coordinate = re.fullmatch(
                                    r"([A-D])([1-9]\d{0,4})", node.get("r", "")
                                )
                                if not coordinate or int(coordinate.group(2)) > 10001:
                                    raise ValueError(
                                        "Solo se permiten columnas A a D y hasta 10.000 movimientos."
                                    )
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("Archivo XLSX inválido o inseguro.") from error
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError("El XLSX debe contener exactamente una hoja.")
        sheet = workbook.active
        sheet.reset_dimensions()
        output = io.StringIO()
        writer = csv.writer(output)
        for index, cells in enumerate(sheet.iter_rows(max_row=10001, max_col=4), start=1):
            values = [cell.value for cell in cells]
            if all(value is None for value in values):
                continue
            if any(cell.data_type in ("f", "e") for cell in cells):
                raise ValueError(f"Fila {index}: reemplaza las fórmulas y errores por valores.")
            if index == 1:
                if values != ["external_id", "date", "amount", "description"]:
                    raise ValueError("Encabezados requeridos: external_id,date,amount,description")
            else:
                if not isinstance(values[0], str) or not isinstance(values[3], str):
                    raise ValueError(f"Fila {index}: ID y descripción deben ser texto.")
                if isinstance(values[1], datetime):
                    if values[1].time() != time():
                        raise ValueError(f"Fila {index}: fecha debe tener solo día, sin hora.")
                    values[1] = values[1].date().isoformat()
                elif isinstance(values[1], date):
                    values[1] = values[1].isoformat()
                if isinstance(values[2], bool):
                    raise ValueError(f"Fila {index}: monto inválido.")
            writer.writerow(values)
        return parse_csv(output.getvalue())
    finally:
        workbook.close()
